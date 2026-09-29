"""Text extraction, chunking, and the background ingestion queue."""

import asyncio
import csv
import io
import json
import logging
import re
from html.parser import HTMLParser
from pathlib import Path
from uuid import UUID

from .config import settings
from .db import pool
from .embeddings import embed_passages

log = logging.getLogger("docvault.ingest")

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".xml", ".html", ".htm", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rb",
    ".c", ".h", ".cpp", ".cs", ".sql", ".sh", ".ini", ".toml", ".cfg", ".conf",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}


class UnsupportedFile(Exception):
    pass


# A section is (page number or None, text).
Section = tuple[int | None, str]


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_sections(path: Path, filename: str) -> list[Section]:
    ext = Path(filename).suffix.lower()
    data = path.read_bytes()

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]

    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [(None, "\n".join(parts))]

    if ext in TEXT_EXTENSIONS:
        text = _decode(data)
        if ext in (".html", ".htm"):
            parser = _HTMLText()
            parser.feed(text)
            text = "".join(parser.parts)
        elif ext == ".json":
            try:
                text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
            except ValueError:
                pass
        elif ext in (".csv", ".tsv"):
            rows = csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else ",")
            text = "\n".join(" | ".join(r) for r in rows)
        return [(None, text)]

    raise UnsupportedFile(f"Unsupported file type: {ext or 'no extension'}")


def chunk_sections(sections: list[Section]) -> list[Section]:
    """Split each section into overlapping chunks, preferring paragraph then sentence boundaries."""
    size, overlap = settings.chunk_size, settings.chunk_overlap
    chunks: list[Section] = []
    for page, text in sections:
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                window = text[start:end]
                for sep in ("\n\n", "\n", ". ", " "):
                    cut = window.rfind(sep)
                    if cut > size // 2:
                        end = start + cut + len(sep)
                        break
            piece = text[start:end].strip()
            if piece:
                chunks.append((page, piece))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return chunks


# Embedding is CPU-heavy on the NAS, so documents are ingested one at a time.
_queue: asyncio.Queue[UUID] = asyncio.Queue()


def enqueue(doc_id: UUID):
    _queue.put_nowait(doc_id)


async def worker():
    while True:
        doc_id = await _queue.get()
        try:
            await _process(doc_id)
        except Exception as e:  # keep the worker alive whatever happens
            log.exception("Ingestion failed for %s", doc_id)
            await _mark_failed(doc_id, str(e))
        finally:
            _queue.task_done()


async def _mark_failed(doc_id: UUID, error: str):
    async with pool.connection() as conn:
        await conn.execute(
            "UPDATE documents SET status = 'failed', error = %s WHERE id = %s", (error[:2000], doc_id)
        )


async def _process(doc_id: UUID):
    async with pool.connection() as conn:
        doc = await (await conn.execute("SELECT * FROM documents WHERE id = %s", (doc_id,))).fetchone()
    if doc is None:
        return  # deleted while queued

    path = settings.upload_dir / str(doc_id)
    sections = await asyncio.to_thread(extract_sections, path, doc["filename"])
    chunks = chunk_sections(sections)
    if not chunks:
        raise ValueError("No extractable text (scanned PDFs need OCR first)")

    log.info("Embedding %d chunks for %s", len(chunks), doc["filename"])
    vectors = await asyncio.to_thread(embed_passages, [c[1] for c in chunks])

    async with pool.connection() as conn, conn.transaction():
        await conn.execute("DELETE FROM chunks WHERE document_id = %s", (doc_id,))
        async with conn.cursor() as cur:
            await cur.executemany(
                "INSERT INTO chunks (document_id, chunk_index, page, content, embedding)"
                " VALUES (%s, %s, %s, %s, %s)",
                [(doc_id, i, page, text, vec) for i, ((page, text), vec) in enumerate(zip(chunks, vectors))],
            )
        await conn.execute(
            "UPDATE documents SET status = 'ready', error = NULL, chunk_count = %s WHERE id = %s",
            (len(chunks), doc_id),
        )
    log.info("Ingested %s (%d chunks)", doc["filename"], len(chunks))


async def requeue_unfinished():
    """Documents left 'processing' by a restart get picked up again."""
    async with pool.connection() as conn:
        rows = await (
            await conn.execute("SELECT id FROM documents WHERE status = 'processing' ORDER BY created_at")
        ).fetchall()
    for row in rows:
        enqueue(row["id"])
